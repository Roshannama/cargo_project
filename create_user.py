from pathlib import Path
import hashlib
import base64
import secrets
import getpass
from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent
USER_FILE = BASE_DIR / "data" / "users.xlsx"


def make_password_hash(password: str) -> str:
    iterations = 600_000
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac(
        "sha256", password.encode("utf-8"), salt, iterations
    )
    return (
        f"pbkdf2_sha256${iterations}$"
        f"{base64.urlsafe_b64encode(salt).decode()}$"
        f"{base64.urlsafe_b64encode(digest).decode()}"
    )


def main():
    username = input("Username: ").strip()
    password = getpass.getpass("Password: ")

    if not username or not password:
        print("Username and password are required.")
        return

    USER_FILE.parent.mkdir(parents=True, exist_ok=True)

    if USER_FILE.exists():
        wb = load_workbook(USER_FILE)
        ws = wb["Users"]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Users"
        ws.append(["username", "password_hash"])

    existing = {
        str(row[0]).strip()
        for row in ws.iter_rows(min_row=2, values_only=True)
        if row[0]
    }

    if username in existing:
        print("That username already exists.")
        wb.close()
        return

    ws.append([username, make_password_hash(password)])
    wb.save(USER_FILE)
    wb.close()

    print(f"User '{username}' added successfully.")
    print("The password is stored as a hash, not plaintext.")


if __name__ == "__main__":
    main()
