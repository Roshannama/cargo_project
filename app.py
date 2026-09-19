import os
import hashlib
import hmac
import base64
from pathlib import Path

from flask import Flask, render_template, request, jsonify, session, redirect, url_for
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
USER_FILE = BASE_DIR / "data" / "users.xlsx"

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "change-this-secret-key-before-production")


def make_password_hash(password: str) -> str:
    """Create a salted PBKDF2 password hash."""
    iterations = 600_000
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac(
        "sha256", password.encode("utf-8"), salt, iterations
    )
    return (
        f"pbkdf2_sha256${iterations}$"
        f"{base64.urlsafe_b64encode(salt).decode()}$"
        f"{base64.urlsafe_b64encode(digest).decode()}"
    )


def check_password(password: str, stored_hash: str) -> bool:
    """Verify a password against the stored PBKDF2 hash."""
    try:
        algorithm, iterations, salt_b64, digest_b64 = stored_hash.split("$")
        if algorithm != "pbkdf2_sha256":
            return False

        iterations = int(iterations)
        salt = base64.urlsafe_b64decode(salt_b64.encode())
        expected = base64.urlsafe_b64decode(digest_b64.encode())

        actual = hashlib.pbkdf2_hmac(
            "sha256", password.encode("utf-8"), salt, iterations
        )
        return hmac.compare_digest(actual, expected)
    except (ValueError, TypeError):
        return False


def authenticate_user(username: str, password: str) -> bool:
    if not USER_FILE.exists():
        return False

    wb = load_workbook(USER_FILE, read_only=True, data_only=True)
    ws = wb["Users"]

    try:
        for row in ws.iter_rows(min_row=2, values_only=True):
            stored_username, stored_password_hash = row[:2]

            if stored_username == username and stored_password_hash:
                return check_password(password, stored_password_hash)
        return False
    finally:
        wb.close()


@app.get("/")
def home():
    return render_template("index.html")


@app.post("/login")
def login():
    data = request.get_json(silent=True) or {}
    username = str(data.get("username", "")).strip()
    password = str(data.get("password", ""))

    if not username or not password:
        return jsonify({
            "success": False,
            "message": "Enter both username and password."
        }), 400

    if authenticate_user(username, password):
        session["username"] = username
        return jsonify({
            "success": True,
            "redirect": url_for("welcome")
        })

    return jsonify({
        "success": False,
        "message": "Invalid username or password."
    }), 401


@app.get("/welcome")
def welcome():
    if "username" not in session:
        return redirect(url_for("home"))
    return render_template("welcome.html", username=session["username"])


@app.get("/logout")
def logout():
    session.clear()
    return redirect(url_for("home"))


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
